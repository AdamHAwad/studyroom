from pathlib import Path
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.colors import HexColor
from PIL import Image,ImageDraw
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches
p=Path('tests/fixtures')
# Original unlabeled shape diagram, with lecture text below. The image is a meaningful visual cue.
im=Image.new('RGB',(1000,600),'white');d=ImageDraw.Draw(im)
d.polygon([(220,440),(500,100),(780,440)],outline='#4255ff',width=12)
im.save(p/'triangle.png')
text='A triangle has three sides and three vertices. A square has four equal sides and four right angles. A circle has no straight sides or vertices. A regular pentagon has five equal sides. The unlabeled diagram above shows a triangle.'
c=Canvas(str(p/'geometry.pdf'),pagesize=(612,792));c.setTitle('Geometry lecture')
c.setFont('Helvetica-Bold',22);c.drawString(55,735,'Recognizing geometric shapes')
c.drawImage(str(p/'triangle.png'),55,370,width=500,height=300)
c.setFont('Helvetica',12)
for i,line in enumerate(['A triangle has three sides and three vertices.','A square has four equal sides and four right angles.','A circle has no straight sides or vertices.','A regular pentagon has five equal sides.','The unlabeled diagram above shows a triangle.']):c.drawString(55,330-i*22,line)
c.save()
doc=Document();doc.add_heading('Recognizing geometric shapes',0);doc.add_picture(str(p/'triangle.png'),width=Inches(5));doc.add_paragraph(text);doc.save(p/'geometry.docx')
prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);box=slide.shapes.add_textbox(PInches(.5),PInches(.2),PInches(9),PInches(.7));box.text='Recognizing geometric shapes';slide.shapes.add_picture(str(p/'triangle.png'),PInches(2),PInches(1),width=PInches(6));box=slide.shapes.add_textbox(PInches(.5),PInches(5.3),PInches(9),PInches(1.6));box.text=text;prs.save(p/'geometry.pptx')
# Scanned PDF for Apple Vision.
c=Canvas(str(p/'scanned.pdf'),pagesize=(612,792));scan=Image.new('RGB',(1600,1000),'white');dr=ImageDraw.Draw(scan)
from PIL import ImageFont
font=ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc',50)
dr.text((80,150),'A triangle has three sides and three vertices.',font=font,fill='black');scan.save(p/'scan.png');c.drawImage(str(p/'scan.png'),0,200,width=612,height=382);c.save()
